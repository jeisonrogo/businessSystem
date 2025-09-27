"""
Utilidades para exportación de datos a formato Excel (.xlsx)
Funciones para generar archivos Excel con formato profesional
"""

from datetime import datetime
from decimal import Decimal
from typing import List, Dict, Any, Optional
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..domain.models.product import Product
from ..domain.models.movimiento_inventario import MovimientoInventario, TipoMovimiento


class ExcelExporter:
    """Clase para generar exportaciones Excel con formato profesional"""

    # Colores del tema empresarial
    HEADER_COLOR = "366092"  # Azul corporativo
    ALTERNATE_ROW_COLOR = "F8F9FA"  # Gris claro

    def __init__(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active

    def _apply_header_style(self, cell):
        """Aplica estilo a las celdas de encabezado"""
        cell.font = Font(bold=True, color="FFFFFFFF", size=11)
        cell.fill = PatternFill(start_color=self.HEADER_COLOR, end_color=self.HEADER_COLOR, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

    def _apply_data_style(self, cell, is_alternate_row=False):
        """Aplica estilo a las celdas de datos"""
        cell.font = Font(size=10)
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        if is_alternate_row:
            cell.fill = PatternFill(
                start_color=self.ALTERNATE_ROW_COLOR,
                end_color=self.ALTERNATE_ROW_COLOR,
                fill_type="solid"
            )

    def _auto_adjust_columns(self, worksheet: Worksheet):
        """Ajusta automáticamente el ancho de las columnas"""
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)  # Máximo 50 caracteres
            worksheet.column_dimensions[column_letter].width = adjusted_width

    def _format_currency(self, value: Any) -> str:
        """Formatea valores monetarios"""
        if value is None:
            return "$0"

        try:
            if isinstance(value, Decimal):
                amount = float(value)
            else:
                amount = float(value)
            return f"${amount:,.2f}"
        except:
            return "$0"

    def _format_quantity_with_sign(self, tipo_movimiento: TipoMovimiento, cantidad: int) -> str:
        """Formatea cantidad con signo según el tipo de movimiento"""
        if tipo_movimiento in [TipoMovimiento.ENTRADA]:
            return f"+{cantidad}"
        elif tipo_movimiento in [TipoMovimiento.SALIDA, TipoMovimiento.MERMA]:
            return f"-{cantidad}"
        else:
            return str(cantidad)

    def export_kardex_to_excel(self, kardex_data: Dict[str, Any], product: Product) -> BytesIO:
        """
        Exporta datos de kardex a formato Excel
        """
        # Configurar worksheet
        self.worksheet.title = f"Kardex {product.sku}"

        # Información del producto (título)
        self.worksheet.merge_cells('A1:I1')
        title_cell = self.worksheet['A1']
        title_cell.value = f"KARDEX DE PRODUCTO: {product.nombre}"
        title_cell.font = Font(bold=True, size=14, color=self.HEADER_COLOR)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Información del producto (subtítulo)
        self.worksheet.merge_cells('A2:I2')
        subtitle_cell = self.worksheet['A2']
        subtitle_cell.value = f"SKU: {product.sku} | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        subtitle_cell.font = Font(size=10, color="666666")
        subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Información resumida
        row = 4
        summary_data = [
            ["Stock Actual:", kardex_data.get('stock_actual', 0)],
            ["Costo Promedio:", self._format_currency(kardex_data.get('costo_promedio_actual', 0))],
            ["Valor Inventario:", self._format_currency(kardex_data.get('valor_inventario', 0))],
            ["Total Movimientos:", kardex_data.get('total_movimientos', 0)]
        ]

        for i, (label, value) in enumerate(summary_data):
            label_cell = self.worksheet.cell(row=row, column=i*2+1)
            value_cell = self.worksheet.cell(row=row, column=i*2+2)

            label_cell.value = label
            label_cell.font = Font(bold=True, size=10)
            value_cell.value = value
            value_cell.font = Font(size=10)

        # Encabezados de la tabla
        headers = [
            "Fecha", "Tipo de Movimiento", "Cantidad", "Precio Unitario",
            "Costo Unitario", "Stock Anterior", "Stock Posterior", "Referencia", "Observaciones"
        ]

        header_row = 6
        for col, header in enumerate(headers, 1):
            cell = self.worksheet.cell(row=header_row, column=col)
            cell.value = header
            self._apply_header_style(cell)

        # Datos de movimientos
        movimientos = kardex_data.get('movimientos', [])
        type_labels = {
            TipoMovimiento.ENTRADA: "Entrada",
            TipoMovimiento.SALIDA: "Salida",
            TipoMovimiento.MERMA: "Merma",
            TipoMovimiento.AJUSTE: "Ajuste"
        }

        for row_idx, movimiento in enumerate(movimientos, header_row + 1):
            is_alternate = (row_idx - header_row) % 2 == 0

            # Datos del movimiento
            data_row = [
                datetime.fromisoformat(movimiento['created_at'].replace('Z', '+00:00')).strftime('%d/%m/%Y %H:%M'),
                type_labels.get(movimiento['tipo_movimiento'], movimiento['tipo_movimiento']),
                self._format_quantity_with_sign(movimiento['tipo_movimiento'], movimiento['cantidad']),
                self._format_currency(movimiento.get('precio_unitario')),
                self._format_currency(movimiento.get('costo_unitario')),
                movimiento['stock_anterior'],
                movimiento['stock_posterior'],
                movimiento.get('referencia', ''),
                movimiento.get('observaciones', '')
            ]

            for col, value in enumerate(data_row, 1):
                cell = self.worksheet.cell(row=row_idx, column=col)
                cell.value = value
                self._apply_data_style(cell, is_alternate)

        # Ajustar columnas
        self._auto_adjust_columns(self.worksheet)

        # Generar archivo en memoria
        excel_file = BytesIO()
        self.workbook.save(excel_file)
        excel_file.seek(0)

        return excel_file

    def export_movements_to_excel(self, movements: List[Dict[str, Any]]) -> BytesIO:
        """
        Exporta movimientos de inventario a formato Excel
        """
        # Configurar worksheet
        self.worksheet.title = "Movimientos Inventario"

        # Título
        self.worksheet.merge_cells('A1:M1')
        title_cell = self.worksheet['A1']
        title_cell.value = "MOVIMIENTOS DE INVENTARIO"
        title_cell.font = Font(bold=True, size=14, color=self.HEADER_COLOR)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Subtítulo
        self.worksheet.merge_cells('A2:M2')
        subtitle_cell = self.worksheet['A2']
        subtitle_cell.value = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        subtitle_cell.font = Font(size=10, color="666666")
        subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Encabezados
        headers = [
            "ID", "Fecha", "Producto", "SKU", "Tipo de Movimiento", "Cantidad",
            "Precio Unitario", "Costo Unitario", "Valor Total", "Stock Anterior",
            "Stock Posterior", "Referencia", "Observaciones"
        ]

        header_row = 4
        for col, header in enumerate(headers, 1):
            cell = self.worksheet.cell(row=header_row, column=col)
            cell.value = header
            self._apply_header_style(cell)

        # Datos de movimientos
        type_labels = {
            TipoMovimiento.ENTRADA: "Entrada",
            TipoMovimiento.SALIDA: "Salida",
            TipoMovimiento.MERMA: "Merma",
            TipoMovimiento.AJUSTE: "Ajuste"
        }

        for row_idx, movement in enumerate(movements, header_row + 1):
            is_alternate = (row_idx - header_row) % 2 == 0

            # Calcular valor total
            precio = movement.get('precio_unitario', 0)
            cantidad = movement.get('cantidad', 0)
            valor_total = float(precio) * cantidad if precio else 0

            data_row = [
                movement.get('id', ''),
                datetime.fromisoformat(movement['created_at'].replace('Z', '+00:00')).strftime('%d/%m/%Y %H:%M'),
                movement.get('producto', {}).get('nombre', 'N/A'),
                movement.get('producto', {}).get('sku', 'N/A'),
                type_labels.get(movement['tipo_movimiento'], movement['tipo_movimiento']),
                self._format_quantity_with_sign(movement['tipo_movimiento'], movement['cantidad']),
                self._format_currency(movement.get('precio_unitario')),
                self._format_currency(movement.get('costo_unitario')),
                self._format_currency(valor_total),
                movement['stock_anterior'],
                movement['stock_posterior'],
                movement.get('referencia', ''),
                movement.get('observaciones', '')
            ]

            for col, value in enumerate(data_row, 1):
                cell = self.worksheet.cell(row=row_idx, column=col)
                cell.value = value
                self._apply_data_style(cell, is_alternate)

        # Ajustar columnas
        self._auto_adjust_columns(self.worksheet)

        # Generar archivo en memoria
        excel_file = BytesIO()
        self.workbook.save(excel_file)
        excel_file.seek(0)

        return excel_file


def create_excel_response(excel_file: BytesIO, filename: str):
    """
    Crea una respuesta HTTP con el archivo Excel
    """
    from fastapi.responses import StreamingResponse

    return StreamingResponse(
        BytesIO(excel_file.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )